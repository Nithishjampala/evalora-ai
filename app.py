
import os
import re
import json
import time
import shutil
import subprocess
import tempfile
from pathlib import Path

import gradio as gr
import pymupdf
import docx
import gspread

from pptx import Presentation
from striprtf.striprtf import rtf_to_text
from dotenv import load_dotenv
from google import genai
from google.genai import types
from google.oauth2.service_account import Credentials


# =====================================================
# 1. CONFIGURATION
# =====================================================

BASE_DIR = Path(__file__).resolve().parent

load_dotenv(BASE_DIR / ".env")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

SERVICE_ACCOUNT_FILE = Path(
    os.getenv(
        "GOOGLE_SERVICE_ACCOUNT_FILE",
        str(BASE_DIR / "service_account.json")
    )
)

SPREADSHEET_ID = (
    "1L8FuY_CFURGrRkWbEFAG_QsHTyF0oUJYfpCfNAl-sOw"
)

WORKSHEET_NAME = os.getenv(
    "GOOGLE_WORKSHEET_NAME",
    "Evaluations"
)

# Keep your configured model names unchanged.
MODELS = [
    "gemini-3.8-flash",
    "gemini-3.6-flash",
    "gemini-3.5-flash-lite",
]

RATING_SCORES = {
    "Outstanding": 100,
    "Very Good": 85,
    "Good": 70,
    "Needs Improvement": 40,
    "Incomplete": 0,
}

# Exactly nine Google Sheets columns.
SHEET_HEADERS = [
    "Name",
    "Project ID",
    "Project Name",
    "Phase Type",
    "Questions",
    "AI Confidence Score",
    "Feedback",
    "Suggestions",
    "Overall Score",
]

SUPPORTED_FILES = [
    ".txt",
    ".pdf",
    ".doc",
    ".docx",
    ".ppt",
    ".pptx",
    ".rtf",
]


# =====================================================
# 2. PHASE QUESTIONS
# =====================================================

PHASES = {
    "Ideation": [
        "How well did the team articulate the problem statement and support it with relevant statistics, data or evidence?",
        "Is the proposed solution addressing a meaningful need for the identified users/stakeholders?",
        "How well has the team researched existing solutions and competitors?",
        "Is the proposed solution feasible and unique?",
        "How would you rate the overall quality of the Ideation submission?",
        "Additional feedback to candidate",
    ],

    "Design": [
        "How effectively does the design demonstrate the proposed solution?",
        "How well does the design address the identified problem?",
        "How effectively has the team demonstrated the core features?",
        "How feasible is the design from a technical perspective?",
        "How would you rate the overall quality of the Design submission?",
        "Additional feedback to candidate",
    ],

    "Build": [
        "How effectively has the team implemented the proposed solution?",
        "How well does the implementation address the identified problem?",
        "How effectively has the team demonstrated the core features?",
        "How feasible is the implementation from a technical perspective?",
        "How would you rate the overall quality of the Build submission?",
        "Additional feedback to candidate",
    ],

    "Demo": [
        "How effectively does the final solution address the stated problem?",
        "How well does the solution demonstrate its key features?",
        "How effectively has the team demonstrated testing and validation?",
        "How feasible and sustainable is the proposed solution?",
        "How would you rate the overall quality of the final submission?",
        "Additional feedback to candidate",
    ],
}


# =====================================================
# 3. DOCUMENT EXTRACTION
# =====================================================

def extract_document(file_path):

    if not file_path:
        raise ValueError("Please upload a document.")

    file_path = str(file_path)
    extension = Path(file_path).suffix.lower()

    text = ""

    # TXT
    if extension == ".txt":

        for encoding in (
            "utf-8-sig",
            "utf-8",
            "cp1252"
        ):
            try:
                text = Path(file_path).read_text(
                    encoding=encoding
                )
                break

            except UnicodeDecodeError:
                continue

    # PDF
    elif extension == ".pdf":

        with pymupdf.open(file_path) as pdf:
            text = "\n".join(
                page.get_text()
                for page in pdf
            )

    # DOCX
    elif extension == ".docx":

        document = docx.Document(file_path)

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:

            for row in table.rows:

                parts.append(
                    " | ".join(
                        cell.text
                        for cell in row.cells
                    )
                )

        text = "\n".join(parts)

    # PPTX
    elif extension == ".pptx":

        presentation = Presentation(file_path)

        slide_text = []

        for number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            parts = [f"Slide {number}"]

            for shape in slide.shapes:

                if getattr(
                    shape,
                    "has_text_frame",
                    False
                ):

                    value = shape.text.strip()

                    if value:
                        parts.append(value)

                if getattr(
                    shape,
                    "has_table",
                    False
                ):

                    for row in shape.table.rows:

                        parts.append(
                            " | ".join(
                                cell.text
                                for cell in row.cells
                            )
                        )

            # Speaker notes
            try:

                if slide.has_notes_slide:

                    notes = (
                        slide.notes_slide
                        .notes_text_frame.text
                    )

                    if notes.strip():

                        parts.append(
                            "Speaker notes: " + notes
                        )

            except Exception:
                pass

            slide_text.append(
                "\n".join(parts)
            )

        text = "\n\n".join(slide_text)

    # RTF
    elif extension == ".rtf":

        raw_text = Path(file_path).read_text(
            encoding="utf-8",
            errors="replace"
        )

        text = rtf_to_text(raw_text)

    # Legacy DOC and PPT
    elif extension in (".doc", ".ppt"):

        if not shutil.which("libreoffice"):

            raise ValueError(
                f"{extension} files require LibreOffice. "
                "Convert the file to DOCX or PPTX and upload again."
            )

        target_extension = (
            ".docx"
            if extension == ".doc"
            else ".pptx"
        )

        target_format = (
            "docx"
            if extension == ".doc"
            else "pptx"
        )

        with tempfile.TemporaryDirectory(
            prefix="evalora_convert_"
        ) as temp_dir:

            process = subprocess.run(
                [
                    "libreoffice",
                    "--headless",
                    "--convert-to",
                    target_format,
                    "--outdir",
                    temp_dir,
                    file_path,
                ],
                capture_output=True,
                text=True,
                timeout=90,
            )

            converted_file = (
                Path(temp_dir)
                /
                (
                    Path(file_path).stem
                    + target_extension
                )
            )

            if (
                process.returncode != 0
                or not converted_file.exists()
            ):

                raise ValueError(
                    "Could not convert the Office document. "
                    "Upload a DOCX or PPTX file instead."
                )

            text = extract_document(
                str(converted_file)
            )

    else:

        raise ValueError(
            "Unsupported format. Upload TXT, PDF, DOC, "
            "DOCX, PPT, PPTX or RTF."
        )

    if not text or not text.strip():

        raise ValueError(
            "No readable text was found. Scanned PDFs "
            "and image-only slides may require OCR."
        )

    return text.strip()[:50000]


# =====================================================
# 4. JSON AND RATING HELPERS
# =====================================================

def parse_json(response_text):

    cleaned = (response_text or "").strip()

    cleaned = re.sub(
        r"^```(?:json)?\s*",
        "",
        cleaned,
        flags=re.IGNORECASE
    )

    cleaned = re.sub(
        r"\s*```$",
        "",
        cleaned
    )

    try:

        result = json.loads(cleaned)

    except json.JSONDecodeError:

        match = re.search(
            r"\{.*\}",
            cleaned,
            re.DOTALL
        )

        if not match:

            raise ValueError(
                "Gemini returned invalid JSON."
            )

        result = json.loads(match.group(0))

    if not isinstance(result, dict):

        raise ValueError(
            "Gemini response must be a JSON object."
        )

    return result


def normalize_rating(value):

    value = str(value or "").strip().lower()

    rating_names = [
        ("needs improvement", "Needs Improvement"),
        ("outstanding", "Outstanding"),
        ("very good", "Very Good"),
        ("incomplete", "Incomplete"),
        ("good", "Good"),
    ]

    for phrase, rating in rating_names:

        if re.search(
            r"\b" + re.escape(phrase) + r"\b",
            value
        ):

            return rating

    return "Incomplete"


def normalize_answers(raw_answers, questions):

    answers_lookup = {}

    if isinstance(raw_answers, list):

        for item in raw_answers:

            if not isinstance(item, dict):
                continue

            question = str(
                item.get("question", "")
            ).strip()

            answer = item.get(
                "answer",
                item.get("rating", "")
            )

            if question:
                answers_lookup[question] = answer

    normalized = []

    for index, question in enumerate(questions):

        answer = answers_lookup.get(
            question,
            ""
        )

        if index < 5:

            answer = normalize_rating(answer)

        else:

            answer = str(
                answer or ""
            ).strip()

            if not answer:

                answer = (
                    "No additional feedback provided."
                )

        normalized.append({
            "question": question,
            "answer": answer
        })

    return normalized


# =====================================================
# 5. GUIDELINE COMPLIANCE NORMALIZATION
# =====================================================

def normalize_compliance(raw_compliance):

    if not isinstance(raw_compliance, list):
        return []

    valid_statuses = {
        "satisfied": "Satisfied",
        "partially satisfied": "Partially Satisfied",
        "not satisfied": "Not Satisfied",
        "not applicable": "Not Applicable",
    }

    normalized = []

    for item in raw_compliance:

        if not isinstance(item, dict):
            continue

        requirement = str(
            item.get("requirement", "")
        ).strip()

        if not requirement:
            continue

        raw_status = str(
            item.get("status", "")
        ).strip().lower()

        status = valid_statuses.get(
            raw_status,
            "Not Satisfied"
        )

        evidence = str(
            item.get("evidence", "")
        ).strip()

        gap = str(
            item.get("gap", "")
        ).strip()

        if not evidence:

            evidence = (
                "No supporting evidence identified."
            )

        if not gap:

            gap = (
                "No additional gap explanation provided."
            )

        normalized.append({
            "requirement": requirement,
            "status": status,
            "evidence": evidence,
            "gap": gap
        })

    return normalized


# =====================================================
# 6. GEMINI EVALUATION PROMPT
# =====================================================

def create_prompt(
    name,
    project_id,
    project_name,
    phase,
    guidelines,
    submission
):

    questions = PHASES[phase]

    question_text = "\n".join(
        f"Q{i}: {question}"
        for i, question in enumerate(
            questions,
            start=1
        )
    )

    return f"""
You are Evalora AI, an evidence-based project evaluator.

Perform two separate tasks:
A. Phase-specific project evaluation.
B. Guideline compliance assessment.

STUDENT DETAILS
Name: {name}
Project ID: {project_id}
Project Name: {project_name}
Selected Phase: {phase}

PHASE QUESTIONS
{question_text}

EVALUATION GUIDELINES
{guidelines[:20000]}

STUDENT SUBMISSION
{submission[:30000]}

PART A: PHASE EVALUATION

1. Evaluate Q1-Q5 independently using relevant evidence.
2. Judge the work according to the selected phase: {phase}.
3. Do not invent statistics, research, features, tests or results.
4. Use exactly these ratings for Q1-Q5:
   Outstanding
   Very Good
   Good
   Needs Improvement
   Incomplete
5. Outstanding means strong evidence with no major gaps.
6. Very Good means well addressed with minor gaps.
7. Good means reasonably addressed, but meaningful details are missing.
8. Needs Improvement means major evidence gaps remain.
9. Incomplete means absent or not assessable.
10. Q5 must be judged independently as overall quality.
11. Q6 is concise constructive feedback, not a rating.
12. Feedback summarizes strengths and important gaps.
13. Suggestions must be specific and actionable.
14. ai_confidence_score represents confidence in the readability
    and completeness of supplied evidence, from 0 to 100.
15. Do not return an overall score. The application calculates it.

Rating score mapping:
Outstanding = 100
Very Good = 85
Good = 70
Needs Improvement = 40
Incomplete = 0

PART B: GUIDELINE COMPLIANCE

Identify distinct, actionable requirements, mandatory deliverables,
constraints and criteria stated in the guidelines.

Focus on requirements applicable to {phase}. Do not invent
requirements or treat background information as a requirement.

If guidelines explicitly define different requirements for different
phases, apply the requirements for {phase}. Otherwise, assess the
applicable requirements without inventing phase restrictions.

For EACH requirement, search the student submission for relevant
supporting evidence. Compare meaning, not exact wording. Accept
synonyms and semantically equivalent explanations.

Use these exact statuses:

Satisfied:
Sufficient relevant evidence shows the requirement is met.

Partially Satisfied:
Some relevant evidence exists, but important elements are missing.

Not Satisfied:
The requirement is missing, contradicted, or insufficiently supported.

Not Applicable:
The requirement genuinely does not apply according to the guidelines.

Do not use Not Applicable merely because evidence is missing.

For each requirement, return:
- requirement: concise description of the actual requirement.
- status: one of the four exact statuses.
- evidence: specific evidence from the submission. Quote a short
  phrase or faithfully paraphrase. If absent, explicitly say no
  relevant evidence was identified.
- gap: explain missing, incomplete or contradictory elements.
  For satisfied requirements, use "No major gap identified" when apt.

Do not treat keyword overlap as proof of compliance.
Mentioning a feature does not prove implementation or testing.
A claim of success is not independent proof.
Do not invent citations, statistics, experiments, page numbers or evidence.
Avoid duplicate requirements.
Include every distinct applicable requirement you can identify.
Do not arbitrarily limit the compliance list to five items.

If no identifiable requirements exist, return an empty
guideline_compliance list and explain the limitation in feedback.
Do not invent precise criteria from vague guideline text.

Return valid JSON only. Do not return Markdown fences or text
outside the JSON object.

Use this structure:

{{
  "phase_answers": [
    {{"question": "{questions[0]}", "answer": "Good"}},
    {{"question": "{questions[1]}", "answer": "Very Good"}},
    {{"question": "{questions[2]}", "answer": "Needs Improvement"}},
    {{"question": "{questions[3]}", "answer": "Good"}},
    {{"question": "{questions[4]}", "answer": "Needs Improvement"}},
    {{"question": "{questions[5]}", "answer": "Constructive feedback."}}
  ],
  "ai_confidence_score": 80,
  "feedback": "Brief overall feedback.",
  "suggestions": "Specific actionable suggestions.",
  "guideline_compliance": [
    {{
      "requirement": "Requirement from the actual guidelines",
      "status": "Satisfied",
      "evidence": "Relevant evidence from the submission",
      "gap": "Missing elements or remaining gaps"
    }}
  ]
}}

The example values illustrate the required structure only.
Generate the actual results from the supplied documents.
"""


# =====================================================
# 7. GEMINI API WITH MODEL FALLBACK
# =====================================================

def call_gemini(prompt, progress=None):

    if not GEMINI_API_KEY:

        raise ValueError(
            "GEMINI_API_KEY is missing from Render Environment "
            "Variables or your local .env file."
        )

    client = genai.Client(
        api_key=GEMINI_API_KEY
    )

    errors = []

    for index, model_name in enumerate(MODELS):

        try:

            if progress:

                progress(
                    min(0.4 + index * 0.08, 0.62),
                    desc=f"Trying {model_name}..."
                )

            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    temperature=0.2
                )
            )

            if not response.text:

                raise ValueError(
                    "Gemini returned an empty response."
                )

            return (
                parse_json(response.text),
                model_name
            )

        except Exception as error:

            errors.append(
                f"{model_name}: {error}"
            )

            time.sleep(1)

    raise RuntimeError(
        "All configured Gemini models failed. Verify that "
        "these model IDs are available to your API key.\n\n"
        + "\n".join(errors)
    )


# =====================================================
# 8. GOOGLE SHEETS - EXACTLY 9 COLUMNS
# =====================================================

def connect_google_sheets():

    if not SERVICE_ACCOUNT_FILE.is_file():

        raise FileNotFoundError(
            "Google service-account JSON was not found. "
            "On Render, add service_account.json under Secret Files "
            "and set GOOGLE_SERVICE_ACCOUNT_FILE to "
            "/etc/secrets/service_account.json."
        )

    credentials = Credentials.from_service_account_file(
        str(SERVICE_ACCOUNT_FILE),
        scopes=[
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive"
        ]
    )

    client = gspread.authorize(credentials)

    spreadsheet = client.open_by_key(
        SPREADSHEET_ID
    )

    try:

        worksheet = spreadsheet.worksheet(
            WORKSHEET_NAME
        )

    except gspread.WorksheetNotFound:

        worksheet = spreadsheet.add_worksheet(
            title=WORKSHEET_NAME,
            rows=1000,
            cols=9
        )

    existing_headers = worksheet.row_values(1)

    if not existing_headers:

        worksheet.append_row(
            SHEET_HEADERS,
            value_input_option="RAW"
        )

    elif existing_headers != SHEET_HEADERS:

        raise ValueError(
            "Update row 1 of the Evaluations worksheet "
            "to these exact nine headers:\n"
            + " | ".join(SHEET_HEADERS)
        )

    return worksheet


def append_to_google_sheet(row_data):

    if len(row_data) != 9:

        raise ValueError(
            "Google Sheets expects exactly 9 columns."
        )

    worksheet = connect_google_sheets()

    worksheet.append_row(
        row_data,
        value_input_option="RAW"
    )


# =====================================================
# 9. GUIDELINE COMPLIANCE REPORT - WEBPAGE ONLY
# =====================================================

def build_compliance_report(compliance):

    lines = [
        "## Guideline vs Student Submission",
        "",
        "Requirements are assessed using meaning and evidence, "
        "not literal text matching.",
        ""
    ]

    if not compliance:

        lines.extend([
            "**Compliance assessment unavailable.**",
            "",
            "No usable requirement-level results were returned. "
            "Check whether the guidelines contain identifiable "
            "requirements and try again."
        ])

        return "\n".join(lines)

    counts = {
        "Satisfied": 0,
        "Partially Satisfied": 0,
        "Not Satisfied": 0,
        "Not Applicable": 0
    }

    for index, item in enumerate(
        compliance,
        start=1
    ):

        status = item["status"]
        counts[status] += 1

        lines.extend([
            f"### Requirement {index}: {item['requirement']}",
            "",
            f"**Status:** {status}",
            "",
            f"**Student Evidence:** {item['evidence']}",
            "",
            f"**Gap / Improvement:** {item['gap']}",
            "",
            "---",
            ""
        ])

    total = len(compliance)

    applicable = (
        total - counts["Not Applicable"]
    )

    percentage = (
        round(
            counts["Satisfied"] / applicable * 100,
            2
        )
        if applicable > 0
        else None
    )

    lines.extend([
        "### Compliance Summary",
        "",
        f"- Total requirements identified: {total}",
        f"- Satisfied: {counts['Satisfied']}",
        f"- Partially Satisfied: {counts['Partially Satisfied']}",
        f"- Not Satisfied: {counts['Not Satisfied']}",
        f"- Not Applicable: {counts['Not Applicable']}",
        ""
    ])

    if percentage is not None:

        lines.append(
            f"- Fully satisfied applicable requirements: "
            f"{percentage}%"
        )

    else:

        lines.append(
            "No applicable requirements were available "
            "for percentage calculation."
        )

    lines.extend([
        "",
        "Note: This percentage counts only fully satisfied "
        "applicable requirements. It is not the overall project score."
    ])

    return "\n".join(lines)


# =====================================================
# 10. EVALUATION
# =====================================================

def evaluate(
    name,
    project_id,
    project_name,
    phase,
    guideline_text,
    guideline_file,
    submission_file,
    progress=gr.Progress()
):

    try:

        if not all([
            name and name.strip(),
            project_id and project_id.strip(),
            project_name and project_name.strip(),
            phase,
            submission_file
        ]):

            raise ValueError(
                "Enter student name, project ID, project name "
                "and phase, then upload the student submission."
            )

        if phase not in PHASES:

            raise ValueError(
                "Select a valid phase."
            )

        progress(
            0.1,
            desc="Reading documents..."
        )

        # Read uploaded guidelines, if provided.
        uploaded_guidelines = ""

        if guideline_file:

            uploaded_guidelines = extract_document(
                guideline_file
            )

        # Read manually entered guidelines, if provided.
        typed_guidelines = str(
            guideline_text or ""
        ).strip()

        # Combine both sources if both are provided.
        if typed_guidelines and uploaded_guidelines:

            guidelines = (
                "GUIDELINES ENTERED BY THE USER:\n"
                + typed_guidelines
                + "\n\nGUIDELINES FROM THE UPLOADED DOCUMENT:\n"
                + uploaded_guidelines
            )

        elif typed_guidelines:

            guidelines = typed_guidelines

        elif uploaded_guidelines:

            guidelines = uploaded_guidelines

        else:

            raise ValueError(
                "Enter evaluation guidelines or upload a "
                "guidelines document."
            )

        # Student submission is required.
        submission = extract_document(
            submission_file
        )

        progress(
            0.3,
            desc="Preparing phase and compliance evaluation..."
        )

        prompt = create_prompt(
            name.strip(),
            project_id.strip(),
            project_name.strip(),
            phase,
            guidelines,
            submission
        )

        result, used_model = call_gemini(
            prompt,
            progress=progress
        )

        progress(
            0.7,
            desc="Validating ratings and compliance..."
        )

        answers = normalize_answers(
            result.get("phase_answers", []),
            PHASES[phase]
        )

        compliance = normalize_compliance(
            result.get("guideline_compliance", [])
        )

        # Overall score is the average of Q1-Q5.
        overall_score = round(
            sum(
                RATING_SCORES[item["answer"]]
                for item in answers[:5]
            ) / 5,
            2
        )

        try:

            confidence = float(
                result.get(
                    "ai_confidence_score",
                    0
                )
            )

        except (ValueError, TypeError):

            confidence = 0

        confidence = max(
            0,
            min(100, confidence)
        )

        feedback = str(
            result.get("feedback", "")
        ).strip()

        if not feedback:

            feedback = (
                "No overall feedback provided."
            )

        suggestions = str(
            result.get("suggestions", "")
        ).strip()

        if not suggestions:

            suggestions = (
                "No suggestions provided."
            )

        # Prepare Questions column.
        # Q1-Q5 contain their rating.
        # Q6 contains additional candidate feedback.
        questions_text = "\n\n".join(
            (
                f"Q{i}: {item['question']}\n"
                f"Rating: {item['answer']}"
                if i <= 5
                else (
                    f"Q{i}: {item['question']}\n"
                    f"Feedback: {item['answer']}"
                )
            )
            for i, item in enumerate(
                answers,
                start=1
            )
        )

        # Exactly nine columns.
        # Guideline compliance is never saved to Sheets.
        row_data = [
            name.strip(),
            project_id.strip(),
            project_name.strip(),
            phase,
            questions_text,
            confidence,
            feedback,
            suggestions,
            overall_score
        ]

        progress(
            0.82,
            desc="Saving evaluation to Google Sheets..."
        )

        append_to_google_sheet(
            row_data
        )

        # Display results on the webpage.
        ui_lines = [
            "# Evalora AI — Evaluation Results",
            "",
            f"**Student:** {name.strip()}",
            f"**Project ID:** {project_id.strip()}",
            f"**Project:** {project_name.strip()}",
            f"**Phase:** {phase}",
            "",
            "## Phase Ratings",
            ""
        ]

        for index, item in enumerate(
            answers[:5],
            start=1
        ):

            ui_lines.extend([
                f"### Q{index}: {item['question']}",
                "",
                f"**Rating:** {item['answer']}",
                ""
            ])

        ui_lines.extend([
            "## Additional Feedback to Candidate",
            "",
            answers[5]["answer"],
            "",
            "## Overall Feedback",
            "",
            feedback,
            "",
            "## Suggestions",
            "",
            suggestions,
            "",
            build_compliance_report(compliance)
        ])

        progress(
            1.0,
            desc="Evaluation completed."
        )

        status = (
            f"Evaluation completed using {used_model}. "
            "Phase ratings and overall score saved to Google Sheets. "
            "Guideline compliance is displayed on the webpage only."
        )

        return (
            "\n".join(ui_lines),
            status
        )

    except Exception as error:

        message = f"Evaluation failed: {error}"

        return (
            f"## Error\n\n{message}",
            message
        )


# =====================================================
# 11. GRADIO USER INTERFACE
# =====================================================

with gr.Blocks(
    title="Evalora AI",
    theme=gr.themes.Soft()
) as demo:

    gr.Markdown(
        """
        # Evalora AI
        ### AI-powered project evaluation and guideline compliance

        Enter the evaluation guidelines or upload a guidelines document.
        Upload the student submission separately.
        """
    )

    with gr.Row():

        name_input = gr.Textbox(
            label="Student Name",
            placeholder="Enter student name"
        )

        project_id_input = gr.Textbox(
            label="Project ID",
            placeholder="Enter project ID"
        )

    project_name_input = gr.Textbox(
        label="Project Name",
        placeholder="Enter project name"
    )

    phase_input = gr.Dropdown(
        choices=list(PHASES.keys()),
        value="Ideation",
        label="Phase Type"
    )

    gr.Markdown(
        "### Evaluation Guidelines"
    )

    guideline_text_input = gr.Textbox(
        label="Enter Guidelines (optional if a file is uploaded)",
        placeholder=(
            "Paste project requirements, judging criteria "
            "or evaluation guidelines here..."
        ),
        lines=7
    )

    guideline_file_input = gr.File(
        label="Or Upload Guidelines Document",
        file_types=SUPPORTED_FILES,
        type="filepath"
    )

    gr.Markdown(
        "### Student Submission"
    )

    submission_input = gr.File(
        label="Upload Student Submission",
        file_types=SUPPORTED_FILES,
        type="filepath"
    )

    evaluate_button = gr.Button(
        "Evaluate Project",
        variant="primary"
    )

    result_output = gr.Markdown(
        label="Evaluation Results"
    )

    status_output = gr.Textbox(
        label="Status",
        interactive=False
    )

    evaluate_button.click(
        fn=evaluate,
        inputs=[
            name_input,
            project_id_input,
            project_name_input,
            phase_input,
            guideline_text_input,
            guideline_file_input,
            submission_input
        ],
        outputs=[
            result_output,
            status_output
        ]
    )


# =====================================================
# 12. APPLICATION ENTRY POINT - RENDER
# =====================================================

if __name__ == "__main__":

    demo.queue().launch(
        server_name="0.0.0.0",
        server_port=int(
            os.environ.get("PORT", "10000")
        ),
        share=False,
        debug=False
    )
